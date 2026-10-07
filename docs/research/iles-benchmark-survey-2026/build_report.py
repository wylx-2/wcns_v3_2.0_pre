"""Build the editable report from the checked Markdown source.

Uses the Codex bundled Python runtime. No solver files are touched.
"""
from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
STEM = 'ILES湍流验证算例调研报告'
md = (ROOT / (STEM+'.md')).read_text(encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.75), Inches(.7)
section.left_margin, section.right_margin = Inches(.85), Inches(.85)
section.header_distance, section.footer_distance = Inches(.3), Inches(.3)

def font(style, size, cn='宋体', latin='Times New Roman', bold=False):
 style.font.name=latin; style.font.size=Pt(size)
 style.font.bold=bold; style.font.color.rgb=RGBColor(0,0,0)
 style.font.italic=False;style.font.underline=False
 pr=style.element.get_or_add_rPr()
 rf=pr.find(qn('w:rFonts'))
 if rf is None: rf=OxmlElement('w:rFonts'); pr.append(rf)
 rf.set(qn('w:eastAsia'),cn)
 for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  rf.attrib.pop(qn('w:'+key),None)

font(doc.styles['Normal'],11)
pf=doc.styles['Normal'].paragraph_format
pf.line_spacing=1.2; pf.space_after=Pt(6); pf.widow_control=True
for n,s in [('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
 font(doc.styles[n],s,'微软雅黑','Arial',n!='Subtitle')
 p=doc.styles[n].paragraph_format
 p.space_before=Pt(13 if n.startswith('Heading') else 0)
 p.space_after=Pt(7)
 p.keep_with_next=True
for n in ['Header','Footer']:
 font(doc.styles[n],9,'宋体','Times New Roman')
# Remove built-in title paragraph borders inherited from the bundled template.
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.core_properties.title='粗网格 ILES 湍流验证算例调研报告'
doc.core_properties.subject='19 类经典湍流算例的配置、参考数据和验证方法'
doc.core_properties.author='WCNS 项目调研'
doc.core_properties.keywords='ILES, DNS, turbulence, WCNS, benchmark'

# Simple inline notation uses native run sub/superscripts where Unicode has
# small phonetic letters instead of mathematical subscripts.
special={'Re꜀':('Re','c'),'Reᴅ':('Re','D')}
def run_text(p,text,bold=False):
 parts=re.split('(Re꜀|Reᴅ)',text)
 for part in parts:
  if part in special:
   a,b=special[part]; r=p.add_run(a);r.bold=bold
   r=p.add_run(b);r.font.subscript=True;r.bold=bold
  elif part:
   r=p.add_run(part);r.bold=bold

def hyperlink(p,label,url):
 if not re.match(r'^https?://',url): url=(ROOT/url).resolve().as_uri()
 rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 col=OxmlElement('w:color');col.set(qn('w:val'),'205580');pr.append(col)
 rf=OxmlElement('w:rFonts');rf.set(qn('w:eastAsia'),'宋体');rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman');pr.append(rf)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

token=re.compile(r'(\*\*.*?\*\*|\[[^\]]+\]\([^\)]+\))')
def inline(p,text):
 for chunk in token.split(text):
  if chunk.startswith('**') and chunk.endswith('**'):run_text(p,chunk[2:-2],True)
  elif chunk.startswith('['):
   m=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',chunk)
   if m: hyperlink(p,m[1],m[2])
   else:run_text(p,chunk)
  else:run_text(p,chunk)

def table(rows):
 columns=len(rows[0]); t=doc.add_table(rows=1,cols=columns)
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if columns==5 and rows[0][0]=='编号': widths=[.48,1.25,1.55,1.9,1.52]
 elif columns==5: widths=[1.15,1.4,1.4,1.25,1.5]
 elif columns==4: widths=[.5,1.55,2.25,2.4]
 else: widths=[1.75,1.85,3.10]
 if columns==4 and rows[0][0]=='无量纲时间': widths=[1.35,1.8,1.8,1.75]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 pr.append(borders)
 mar=OxmlElement('w:tblCellMar')
 for edge,v in [('top','75'),('bottom','75'),('left','90'),('right','90')]:
  e=OxmlElement('w:'+edge);e.set(qn('w:w'),v);e.set(qn('w:type'),'dxa');mar.append(e)
 pr.append(mar)
 for idx,row in enumerate(rows):
  cells=t.rows[0].cells if idx==0 else t.add_row().cells
  trpr=t.rows[idx]._tr.get_or_add_trPr()
  no=OxmlElement('w:cantSplit');trpr.append(no)
  if idx==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for j,(cell,value) in enumerate(zip(cells,row)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.1
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (j==0 or rows[0][0]=='无量纲时间') else WD_ALIGN_PARAGRAPH.LEFT
   inline(p,value)
   for r in p.runs:r.font.size=Pt(10);r.bold=idx==0
   if idx==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DEE7EF');cell._tc.get_or_add_tcPr().append(sh)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)

hp=section.header.paragraphs[0];hp.text='WCNS　粗网格 ILES 湍流验证算例调研'
hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp=section.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fp.add_run('第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
fp.add_run(' 页')

lines=md.splitlines();i=0
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[-:]+',c) for c in cells): rows.append(cells)
   i+=1
  table(rows);continue
 if line.startswith('# '):
  doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):
  title=line[3:]
  p=doc.add_paragraph(title,'Heading 1')
  # A fresh page for substantial chapters, leaving the opening recommendation
  # on the title page and letting individual case entries flow naturally.
  if title.startswith(('2 ','3 ','11 ')):p.paragraph_format.page_break_before=True
 elif line.startswith('### '):
  title=re.sub(r'^(9)\.(\d)',r'\1 \2',line[4:])
  doc.add_paragraph(title,'Heading 2')
 elif line=='研究背景 参考配置 可信数据与验证方法':
  doc.add_paragraph(line,'Subtitle')
 elif line.startswith('调研日期'):
  p=doc.add_paragraph(line);p.paragraph_format.space_after=Pt(16)
  for r in p.runs:r.font.size=Pt(10)
 else:
  p=doc.add_paragraph();inline(p,line)
  if re.match(r'^R\d{2}　',line):
   p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(8)
   for r in p.runs:r.font.size=Pt(10.5)
 i+=1
out=ROOT/(STEM+'.docx');doc.save(out)
summary={'docx':str(out),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'bytes':out.stat().st_size,'md_sha256':hashlib.sha256(md.encode()).hexdigest()}
(ROOT/'qa'/'build_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=True))
